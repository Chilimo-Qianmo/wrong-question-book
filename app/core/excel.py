# -*- coding: utf-8 -*-
"""答题表（xlsx）解析 —— 移植自旧版 生成错题集.py 并按 E9 加固。

修复的三个旧 bug：
  1. 表头识别只认纯数字：现兼容 1 / "1" / 1.0 / "1.0" / "第1题" / "1题" / "T1" /
     全角数字（１、第１题）等形态（NFKC 归一化后再解析）；表头不在第 1 行时自动定位。
  2. 单元格判定过窄（旧代码只认"数值型且等于 0"，字符串 "0" 漏判、只读第一列）：
     现 0 / 0.0 / "0" / "0.0" / "0分" / 全角零 都算答错；空白/None 不算；
     其它非空非零值不算答错，但会在诊断 notes 里给出取值样例。
  3. 姓名列写死第 0 列：现在默认第 0 列，可用 name_column 指定；
     表头中出现「姓名/学生/名字/名单/考生」时自动定位该列。

命名约定：诊断信息字段与 app/models.py 的 ExcelPeek 对齐
（path / sheet / students / question_columns / missing_columns / class_name /
name_column / notes）。
"""
from __future__ import annotations

import os
import re
import unicodedata
from typing import Any, Iterable, Optional, Sequence

from openpyxl import load_workbook

from .archive import class_from_excel

__all__ = ["load_wrong_answers", "peek_excel", "parse_qno_header",
           "is_wrong_cell", "find_header_row", "detect_name_column"]

DEFAULT_NAME_COLUMN = 0
NAME_HEADER_HINTS = ("姓名", "学生", "名字", "名单", "考生")
MAX_SAMPLES = 8
MAX_HEADER_SAMPLES = 6

_ZERO_RE = re.compile(r"^[+-]?0+(?:\.0+)?$")
_QNO_PATTERNS = (
    re.compile(r"^[Tt]\s*(\d{1,3})$"),                    # T1
    re.compile(r"^第\s*(\d{1,3})\s*题?$"),                # 第1题 / 第 1 题
    re.compile(r"^(\d{1,3})\s*题$"),                      # 1题
)
_QNO_NUM_RE = re.compile(r"^(\d{1,3})(?:\.(\d+))?[\.、,:：]?$")   # 1 / "1" / 1.0 / 1、


# --------------------------------------------------------------------------- #
# 归一化与判定
# --------------------------------------------------------------------------- #
def _norm_text(value) -> str:
    """全角→半角（NFKC）+ 去首尾空白（'１'→'1'、'第１题'→'第1题'）。"""
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def parse_qno_header(cell) -> Optional[int]:
    """把表头单元格解析成题号；不是题号列则返回 None。

    支持：1 / "1" / 1.0 / "1.0" / "1." / "1、" / "第1题" / "1题" / "T1" / 全角数字。
    """
    if cell is None or isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        try:
            val = float(cell)
        except (TypeError, ValueError):
            return None
        if val.is_integer() and 0 < int(val) < 1000:
            return int(val)
        return None

    s = _norm_text(cell)
    if not s:
        return None
    qn = None
    for pat in _QNO_PATTERNS:
        m = pat.match(s)
        if m:
            qn = int(m.group(1))
            break
    if qn is None:
        m = _QNO_NUM_RE.match(s)
        if not m:
            return None
        frac = m.group(2)
        if frac and set(frac) - {"0"}:      # 1.5 之类不是题号列
            return None
        qn = int(m.group(1))
    return qn if 0 < qn < 1000 else None


def _fmt_number(value) -> str:
    try:
        num = float(value)
    except (TypeError, ValueError):
        return str(value)
    if num.is_integer():
        return str(int(num))
    return ("%g" % num)


def is_wrong_cell(value):
    """判断单元格是否表示「答错」。

    返回 (是否答错, 是否空白, 取值样例)：
      * None / '' / 全空白            -> (False, True,  None)   空白不算答错
      * 0 / 0.0 / "0" / "0分" / "0.0" -> (True,  False, None)   答错
      * 其它非空非零值（1、"√"、0.5…）-> (False, False, "样例")  不算答错，但记录样例
    """
    if value is None:
        return False, True, None
    if isinstance(value, bool):
        return False, False, str(value)
    if isinstance(value, (int, float)):
        try:
            num = float(value)
        except (TypeError, ValueError):
            return False, False, str(value)
        if num == 0.0:
            return True, False, None
        return False, False, _fmt_number(value)

    s = _norm_text(value)
    if not s:
        return False, True, None
    compact = re.sub(r"\s+", "", s)
    trimmed = re.sub(r"分$", "", compact)          # "0分" -> "0"
    if _ZERO_RE.match(trimmed):
        return True, False, None
    return False, False, s


# --------------------------------------------------------------------------- #
# 读取与结构定位
# --------------------------------------------------------------------------- #
def _read_rows(excel_path, sheet=None):
    """读取整张表，返回 (sheet 名, rows)。sheet 可为名称（str）或序号（int，从 1 起）。"""
    wb = load_workbook(str(excel_path), data_only=True, read_only=True)
    try:
        if sheet is None:
            ws = wb.active
        elif isinstance(sheet, int):
            ws = wb.worksheets[sheet - 1] if sheet >= 1 else wb.worksheets[sheet]
        else:
            ws = wb[str(sheet)]
        title = ws.title
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        return title, rows
    finally:
        wb.close()


def find_header_row(rows, scan: int = 8) -> int:
    """在前 scan 行里定位表头行：题号列最多的那一行（严格更优才替换，故首行优先）。"""
    best_idx, best_hits = 0, -1
    for idx, row in enumerate(rows[:max(1, scan)]):
        if not row:
            continue
        hits = sum(1 for cell in row if parse_qno_header(cell) is not None)
        if hits > best_hits:
            best_idx, best_hits = idx, hits
    return best_idx


def detect_name_column(header, name_column=None) -> int:
    """姓名列：显式指定优先；否则按「姓名/学生/名字/名单/考生」自动定位；再回退第 0 列。"""
    if name_column is not None:
        try:
            return int(name_column)
        except (TypeError, ValueError):
            return DEFAULT_NAME_COLUMN
    for i, cell in enumerate(header or ()):
        text = _norm_text(cell)
        if not text:
            continue
        if any(hint in text for hint in NAME_HEADER_HINTS):
            return i
    return DEFAULT_NAME_COLUMN


def _scan_columns(header, name_col):
    """扫描表头，返回 (题号->列索引, 重复题号, 未识别表头样例)。"""
    qidx: dict = {}
    dups: list = []
    unknown: list = []
    for i, cell in enumerate(header or ()):
        if i == name_col:
            continue
        qn = parse_qno_header(cell)
        if qn is None:
            text = _norm_text(cell)
            if text and text not in unknown:
                unknown.append(text)
            continue
        if qn in qidx:
            dups.append((qn, qidx[qn], i))
            continue
        qidx[qn] = i
    return qidx, dups, unknown


def _is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return False


def _sample_text(samples: dict, limit: int = MAX_SAMPLES) -> str:
    if not samples:
        return ""
    items = sorted(samples.items(), key=lambda kv: (-kv[1], kv[0]))[:limit]
    text = "、".join("'%s'(%d次)" % (k, v) for k, v in items)
    if len(samples) > limit:
        text += " 等 %d 种" % len(samples)
    return text


# --------------------------------------------------------------------------- #
# 对外接口
# --------------------------------------------------------------------------- #
def load_wrong_answers(excel_path, qnos, name_column=None, sheet=None):
    """解析答题表，返回 ({学生姓名: [错误题号...]}, 诊断信息 dict)。

    参数
    ----
    excel_path : xlsx 路径
    qnos       : 需要统计的题号序列（试卷识别到的题目）
    name_column: 可选，姓名所在列索引（默认自动识别，兜底第 0 列）
    sheet      : 可选，工作表名或序号（默认活动表）

    诊断信息包含 path / sheet / students / question_columns / missing_columns /
    name_column / class_name / notes / rows / wrong_cells / value_samples。
    """
    want: list = []
    for q in (qnos or []):
        try:
            qn = int(q)
        except (TypeError, ValueError):
            continue
        if qn not in want:
            want.append(qn)
    want.sort()

    diag: dict = {
        "path": os.path.abspath(str(excel_path)),
        "sheet": "",
        "students": 0,
        "rows": 0,
        "name_column": DEFAULT_NAME_COLUMN,
        "question_columns": [],
        "missing_columns": sorted(want),
        "class_name": class_from_excel(excel_path),
        "notes": [],
        "wrong_cells": 0,
        "value_samples": [],
    }
    notes = diag["notes"]

    title, rows = _read_rows(excel_path, sheet=sheet)
    diag["sheet"] = title
    diag["rows"] = len(rows)
    if not rows:
        notes.append("表格为空：没有任何行。")
        return {}, diag

    hidx = find_header_row(rows)
    header = rows[hidx]
    if hidx != 0:
        notes.append("表头不在第 1 行，已自动定位到第 %d 行。" % (hidx + 1))

    name_col = detect_name_column(header, name_column)
    diag["name_column"] = name_col
    qidx, dups, unknown = _scan_columns(header, name_col)
    diag["question_columns"] = sorted(qidx)
    if dups:
        notes.append("表头中题号重复的列（已取最左一列）："
                     + "、".join("第%d题(列%d与列%d)" % (q, a, b) for q, a, b in dups))
    if unknown:
        notes.append("未识别的表头列（已忽略）：" + "、".join(
            "'%s'" % u for u in unknown[:MAX_HEADER_SAMPLES]))

    missing = [q for q in want if q not in qidx]
    diag["missing_columns"] = missing
    if missing:
        notes.append("答题表中缺少题号列：%s（这些题一律不判为错题）。"
                     % "、".join(str(q) for q in missing))
    extra = [q for q in sorted(qidx) if q not in want]
    if extra:
        notes.append("答题表中有、但本次未统计的题号列：%s。"
                     % "、".join(str(q) for q in extra))

    wrong: dict = {}
    samples: dict = {}
    dup_names: list = []
    wrong_cells = 0
    for row in rows[hidx + 1:]:
        if not row:
            continue
        raw_name = row[name_col] if name_col < len(row) else None
        if _is_blank(raw_name):
            continue
        name = _norm_text(raw_name) or str(raw_name).strip()
        bad = []
        for qn in want:
            idx = qidx.get(qn)
            if idx is None or idx >= len(row):
                continue
            is_bad, _blank, sample = is_wrong_cell(row[idx])
            if is_bad:
                bad.append(int(qn))
            elif sample is not None:
                samples[sample] = samples.get(sample, 0) + 1
        wrong_cells += len(bad)
        if name in wrong:
            dup_names.append(name)
            wrong[name] = sorted(set(wrong[name]) | set(bad))   # 同名合并，绝不静默丢人
        else:
            wrong[name] = sorted(set(bad))

    diag["students"] = len(wrong)
    diag["wrong_cells"] = wrong_cells
    diag["value_samples"] = sorted(samples)
    if dup_names:
        notes.append("姓名重复 %d 处（已合并为一条）：%s"
                     % (len(dup_names), "、".join(sorted(set(dup_names))[:MAX_HEADER_SAMPLES])))
    if samples:
        notes.append("非零非空取值样例（这些不计为答错）：" + _sample_text(samples))
    if not wrong:
        notes.append("没有读到任何学生姓名（姓名列 = 第 %d 列）。" % (name_col + 1))
    return wrong, diag


def peek_excel(excel_path, sheet=None, name_column=None) -> dict:
    """只做体检、不返回学生数据：给前端「来源体检」卡片用的摘要。

    返回字段与 app/models.py 的 ExcelPeek 对齐：
    path / sheet / students / question_columns / missing_columns / class_name /
    name_column / notes。
    """
    result = {
        "path": os.path.abspath(str(excel_path)),
        "sheet": "",
        "students": 0,
        "question_columns": [],
        "missing_columns": [],
        "class_name": class_from_excel(excel_path),
        "name_column": DEFAULT_NAME_COLUMN,
        "notes": [],
    }
    notes = result["notes"]
    if not excel_path or not os.path.exists(str(excel_path)):
        notes.append("文件不存在：%s" % excel_path)
        return result

    title, rows = _read_rows(excel_path, sheet=sheet)
    result["sheet"] = title
    if not rows:
        notes.append("表格为空：没有任何行。")
        return result

    hidx = find_header_row(rows)
    header = rows[hidx]
    if hidx != 0:
        notes.append("表头不在第 1 行，已自动定位到第 %d 行。" % (hidx + 1))

    name_col = detect_name_column(header, name_column)
    result["name_column"] = name_col
    qidx, dups, unknown = _scan_columns(header, name_col)
    qcols = sorted(qidx)
    result["question_columns"] = qcols
    if dups:
        notes.append("表头中题号重复的列（已取最左一列）："
                     + "、".join("第%d题(列%d与列%d)" % (q, a, b) for q, a, b in dups))
    if unknown:
        notes.append("未识别的表头列（已忽略）：" + "、".join(
            "'%s'" % u for u in unknown[:MAX_HEADER_SAMPLES]))
    notes.append("识别到 %d 个题号列：%s" % (len(qcols),
                 ("%s..%s" % (qcols[0], qcols[-1])) if len(qcols) > 1
                 else ("%s" % qcols[0] if qcols else "无")))

    # 题号序列的缺口（相对于识别到的最小..最大题号），供前端告警
    missing = []
    if qcols:
        present = set(qcols)
        missing = [q for q in range(qcols[0], qcols[-1] + 1) if q not in present]
    result["missing_columns"] = missing
    if missing:
        notes.append("题号列不连续，缺失：%s。" % "、".join(str(q) for q in missing))

    samples: dict = {}
    students = 0
    names: dict = {}
    for row in rows[hidx + 1:]:
        if not row:
            continue
        raw_name = row[name_col] if name_col < len(row) else None
        if _is_blank(raw_name):
            continue
        students += 1
        name = _norm_text(raw_name)
        names[name] = names.get(name, 0) + 1
        for idx in qidx.values():
            if idx >= len(row):
                continue
            is_bad, _blank, sample = is_wrong_cell(row[idx])
            if not is_bad and sample is not None:
                samples[sample] = samples.get(sample, 0) + 1
    result["students"] = students
    dup = [n for n, c in names.items() if c > 1]
    if dup:
        notes.append("姓名重复：%s。" % "、".join(dup[:MAX_HEADER_SAMPLES]))
    if samples:
        notes.append("非零非空取值样例（这些不计为答错）：" + _sample_text(samples))
    if not students:
        notes.append("没有读到任何学生姓名（姓名列 = 第 %d 列）。" % (name_col + 1))
    return result
